"""Write the run's reference table (.xlsx): title, authors, DOI, APA 7th citation, summary.

Input is a JSON list, one object per paper the deliverable cites or selects:
    [{"doi": "10.1038/s41467-019-13462-1",
      "summary": "One or two sentences on what the paper did and found.",
      "basis": "full text",            # or "abstract"
      "title": "Global urban expansion offsets ..."}]   # optional: the title you expect

Title, authors and the APA citation come from the DOI itself (doi.org content
negotiation: CSL-JSON plus the citeproc "apa" style), so they are never retyped by
hand. If you give an expected title and the DOI resolves to a different paper, the
row is flagged: a wrong DOI is caught here rather than in the reference list.

    python export_reference_table.py papers.json references.xlsx

Exit 0 when every DOI resolved and matched, 2 when any row is flagged (the file is
still written; flagged rows sort first, with the reason in the Check column).
Italics are lost in a spreadsheet cell; the APA text is otherwise final, including
`Article <number>` for article-numbered papers, which citeproc leaves out. Responses are cached next to
the output, so a re-run only fetches what is new or failed.
"""
import difflib
import html
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

MAILTO = os.environ.get("CROSSREF_MAILTO", "")
UA = "deep-science-writer/2.0" + (f" (mailto:{MAILTO})" if MAILTO else "")
CSL = "application/vnd.citationstyles.csl+json"
APA = "text/x-bibliography; style=apa; locale=en-US"


def norm_doi(d):
    return re.sub(r"^(https?://(dx\.)?doi\.org/|doi:)", "", (d or "").strip(), flags=re.I).lower()


def fetch(doi, accept, cache):
    key = cache / (re.sub(r"[^a-z0-9]+", "_", doi) + ("_csl.json" if accept == CSL else "_apa.txt"))
    if key.exists():
        return key.read_text(encoding="utf-8")
    req = urllib.request.Request(f"https://doi.org/{doi}", headers={"Accept": accept, "User-Agent": UA})
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                text = r.read().decode("utf-8")
            key.write_text(text, encoding="utf-8")   # only successes are cached
            return text
        except urllib.error.HTTPError as e:
            if e.code == 404:
                raise RuntimeError("DOI not found (404)")
            if e.code not in (429, 500, 502, 503, 504) or attempt == 2:
                raise RuntimeError(f"HTTP {e.code}")
        except (urllib.error.URLError, TimeoutError) as e:
            if attempt == 2:
                raise RuntimeError(f"network error: {e}")
        time.sleep(2 * (attempt + 1))


def authors_of(csl):
    names = []
    for a in csl.get("author", []):
        if "family" in a:
            names.append(f"{a['family']}, {a.get('given', '')}".strip().rstrip(","))
        elif "literal" in a:
            names.append(a["literal"])
    return "; ".join(names)


def first(v):
    return v[0] if isinstance(v, list) and v else (v or "")


def fix_article_number(apa, csl, p, check):
    """citeproc's apa style drops Crossref's `article-number` and prints an eLocator
    such as 084010 as page 84010. APA 7th wants `, Article <number>` after vol(issue).
    papers.json may settle a case by hand: "article_number": "107737", or
    "single_page": true for a genuine one-page paper."""
    num, page = p.get("article_number") or csl.get("article-number"), str(csl.get("page") or "")
    if not num and page and not re.search(r"[-–]", page) and not p.get("single_page"):
        if page.startswith("0") or re.search(r"[A-Za-z]", page):
            num = page                      # a page cannot start with 0 or hold a letter
        else:
            check.append(f"single page '{page}': set article_number or single_page in papers.json")
    if not num:
        return apa
    vol, iss = str(csl.get("volume") or ""), str(csl.get("issue") or "")
    vi = f"{vol}({iss})" if iss else vol
    n = 0
    if vol:   # anchored on ", vol(issue)" right before the DOI link, so a title or initial is never hit
        apa, n = re.subn(r", " + re.escape(vi) + r"(, [^.,]+)?\.(?=\s+https?://|\s*$)",
                         f", {vi}, Article {num}.", apa, count=1)
    if not n:
        check.append(f"article number {num} not placed: check the APA string")
    return apa


def plain(s):
    """Crossref titles carry inline markup (CO<sub>2</sub>, <i>Genus</i>)."""
    return html.unescape(re.sub(r"<[^>]+>", "", s or ""))


ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")   # openpyxl refuses these (e.g. a PDF form feed)


def similar(a, b):
    clean = lambda s: re.sub(r"[^a-z0-9 ]", "", html.unescape(s).lower())
    return difflib.SequenceMatcher(None, clean(a), clean(b)).ratio()


def main(src, out):
    papers = json.loads(Path(src).read_text(encoding="utf-8"))
    out = Path(out)
    cache = out.parent / ".doi_cache"
    cache.mkdir(parents=True, exist_ok=True)
    rows, seen, flagged = [], set(), 0
    for p in papers:
        doi, check = norm_doi(p.get("doi")), []
        title, authors, apa = p.get("title", ""), p.get("authors", ""), p.get("apa", "")
        if not doi:
            check.append("no DOI: title, authors and APA are as given, not verified")
        elif doi in seen:
            check.append("duplicate DOI in input")
        else:
            seen.add(doi)
            try:
                csl = json.loads(fetch(doi, CSL, cache))
                apa = re.sub(r"\s+", " ", plain(fetch(doi, APA, cache))).strip()
                apa = fix_article_number(apa, csl, p, check)
                resolved = plain(first(csl.get("title")))
                if p.get("title") and similar(p["title"], resolved) < 0.6:
                    check.append(f"TITLE MISMATCH: DOI resolves to '{resolved[:90]}'")
                title, authors = resolved, authors_of(csl)
            except RuntimeError as e:
                check.append(f"DOI did not resolve: {e}")
        if not p.get("summary"):
            check.append("no summary")
        flagged += bool(check)
        rows.append([title, authors, doi, apa, p.get("summary", ""), p.get("basis", ""), "; ".join(check)])
        print(f"{'FLAG' if check else 'ok  '} {doi or '(no doi)'} {'; '.join(check)}")

    rows.sort(key=lambda r: (not r[6], (r[3] or r[0] or "").lower()))   # flagged rows first, then APA order
    wb = Workbook()
    ws = wb.active
    ws.title = "References"
    head = ["Title", "Authors", "DOI", "APA 7th citation", "Summary", "Summary basis", "Check"]
    ws.append(head)
    for r in rows:
        ws.append([ILLEGAL.sub(" ", v) if isinstance(v, str) else v for v in r])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE5F0")
    for col, width in zip("ABCDEFG", (45, 35, 28, 70, 70, 14, 30)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if row[2].value:
            row[2].hyperlink = f"https://doi.org/{row[2].value}"
            row[2].font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out)
    cell = lambda v: re.sub(r"\s+", " ", ILLEGAL.sub(" ", str(v or ""))).replace("|", "/")
    md = [f"# Reference table ({out.name})", "", "| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
    md += ["| " + " | ".join(cell(v) for v in r) + " |" for r in rows]
    out.with_suffix(".md").write_text("\n".join(md) + "\n", encoding="utf-8")   # the RAG cannot index .xlsx
    print(f"{len(rows)} rows, {flagged} flagged -> {out} (+ {out.with_suffix('.md').name})")
    return 2 if flagged else 0


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1], sys.argv[2]))
